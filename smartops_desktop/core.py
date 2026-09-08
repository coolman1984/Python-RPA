from __future__ import annotations

import json
import os
import re
import sqlite3
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4
from urllib.parse import urlsplit
from zipfile import ZipFile, BadZipFile

import yaml
from openpyxl import load_workbook

ACTIONS = ("navigate", "click", "fill", "select", "check", "press", "wait", "download", "validate_xlsx", "nexacro_probe", "demo_export")
DEFAULT_SETTINGS = {
    "cdp_url": "http://127.0.0.1:9222",
    "chrome_launcher": r"D:\WORK\Software Development\GitHub\AI CREW\Mandatory To Use Skills\windows-chrome-launcher\scripts\open_chrome.py",
    "timeout_seconds": 30,
}


def now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def data_root():
    override = os.environ.get("SMARTOPS_DATA_DIR")
    return Path(override) if override else Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "SmartOpsDesktop"


def atomic_text(path: Path, text: str):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=".smartops-", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            stream.write(text)
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)


def http_url(value):
    if not isinstance(value, str):
        raise ValueError("URL must be text.")
    parsed = urlsplit(value)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("Use an http(s) URL without embedded credentials.")
    return value


def cdp_url(value):
    http_url(value)
    if urlsplit(value).hostname not in {"localhost", "127.0.0.1", "::1"}:
        raise ValueError("Chrome connection must use a local loopback address.")
    return value


def validate_workflow(raw):
    if not isinstance(raw, dict) or raw.get("schema_version") != 1:
        raise ValueError("Expected workflow schema_version 1.")
    if not isinstance(raw.get("name"), str) or not raw["name"].strip():
        raise ValueError("Give this workflow a name.")
    if not isinstance(raw.get("steps"), list) or len(raw["steps"]) > 1000:
        raise ValueError("Steps must be a list with at most 1,000 entries.")
    clean = {"schema_version": 1, "id": str(raw.get("id") or uuid4()), "name": raw["name"].strip()[:160], "description": str(raw.get("description", ""))[:2000], "steps": []}
    if not re.fullmatch(r"[a-zA-Z0-9_-]{1,80}", clean["id"]):
        raise ValueError("Invalid workflow ID.")
    for index, step in enumerate(raw["steps"], 1):
        if not isinstance(step, dict) or step.get("action") not in ACTIONS:
            raise ValueError(f"Step {index}: unsupported action.")
        # dict(step) keeps every key the recorder attached: fingerprint, detected_by, capture
        # metadata. Rebuilding a step from known fields only is how that evidence gets lost.
        item = dict(step)
        action = item["action"]
        if action == "navigate":
            http_url(item.get("url", ""))
        if action in {"click", "fill", "select", "check", "press", "download"}:
            if not isinstance(item.get("selector"), str) or not item["selector"].strip():
                raise ValueError(f"Step {index}: a selector is required.")
        if action in {"fill", "select", "press"} and not isinstance(item.get("value"), str):
            raise ValueError(f"Step {index}: a text value is required.")
        if action == "fill" and re.search(r"password|passwd|secret|token|otp", item["selector"], re.I):
            raise ValueError("Credential fields must be filled manually, outside recording.")
        if action == "check" and not isinstance(item.get("checked", True), bool):
            raise ValueError("checked must be true or false.")
        if action == "wait" and (isinstance(item.get("seconds", 1), bool) or not isinstance(item.get("seconds", 1), (float, int)) or not 0 <= item.get("seconds", 1) <= 300):
            raise ValueError("Wait must be between 0 and 300 seconds.")
        if action in {"validate_xlsx", "download"}:
            rows = item.get("min_rows", 1)
            if isinstance(rows, bool) or not isinstance(rows, int) or rows < 0:
                raise ValueError("Minimum rows must be a non-negative integer.")
            if not isinstance(item.get("required_columns", []), list) or any(not isinstance(c, str) for c in item.get("required_columns", [])):
                raise ValueError("Required columns must be a list of names.")
        clean["steps"].append(item)
    return clean


def demo_workflow():
    return {"schema_version": 1, "id": "welcome-demo", "name": "Your first successful run", "description": "A local demo: create a sample report, validate its real Excel contents, and keep the result. No Chrome or sign-in needed.", "steps": [
        {"action": "demo_export", "label": "Create a sample production report"},
        {"action": "validate_xlsx", "label": "Check workbook, rows and columns", "min_rows": 3, "required_columns": ["Date", "Line", "Quantity"]},
    ]}


def validate_xlsx(path, min_rows=1, required_columns=(), sheet=None):
    path = Path(path)
    if not path.is_file():
        raise ValueError("No completed file is available to validate.")
    try:
        with ZipFile(path) as archive:
            names = set(archive.namelist())
            if not {"[Content_Types].xml", "xl/workbook.xml"} <= names:
                raise ValueError("The downloaded bytes are not an OOXML Excel workbook.")
            if sum(x.file_size for x in archive.infolist()) > 512 * 1024 * 1024:
                raise ValueError("Workbook expands beyond the 512 MB validation limit.")
            if archive.testzip():
                raise ValueError("Workbook ZIP integrity check failed.")
        # Passing a stream checks actual contents, independent of the extension.
        with path.open("rb") as stream:
            workbook = load_workbook(stream, read_only=True, data_only=True)
            try:
                if sheet and sheet not in workbook.sheetnames:
                    raise ValueError(f"Worksheet not found: {sheet}")
                selected = workbook[sheet] if sheet else workbook.active
                iterator = selected.iter_rows(values_only=True)
                headers = [str(v).strip() if v is not None else "" for v in next(iterator, ())]
                missing = sorted(set(required_columns) - set(headers))
                if missing:
                    raise ValueError("Missing columns: " + ", ".join(missing))
                count = sum(1 for row in iterator if any(v is not None and v != "" for v in row))
                if count < min_rows:
                    raise ValueError(f"Expected at least {min_rows} data rows; found {count}.")
                return {"format": "OOXML Excel", "sheet": selected.title, "rows": count, "columns": headers, "bytes": path.stat().st_size}
            finally:
                workbook.close()
    except BadZipFile as exc:
        raise ValueError("File is not an XLSX workbook; it may be HTML, legacy XLS, or protected content.") from exc


class Store:
    def __init__(self, root=None):
        self.root = Path(root) if root else data_root()
        self.workflows = self.root / "workflows"
        self.workflows.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(self.root / "smartops.sqlite3")
        self.db.row_factory = sqlite3.Row
        self.db.executescript("""
        CREATE TABLE IF NOT EXISTS workflows (id TEXT PRIMARY KEY, name TEXT NOT NULL, body TEXT NOT NULL, updated TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS runs (id TEXT PRIMARY KEY, workflow_id TEXT, name TEXT, started TEXT, finished TEXT, status TEXT, artifact TEXT, detail TEXT);
        CREATE TABLE IF NOT EXISTS events (id INTEGER PRIMARY KEY, run_id TEXT, time TEXT, body TEXT);
        """)
        self.db.execute("UPDATE runs SET status='interrupted', finished=? WHERE status IN ('running','recording')", (now(),))
        self.db.commit()
        if self.db.execute("SELECT count(*) FROM workflows").fetchone()[0] == 0:
            self.save(demo_workflow())

    def settings(self):
        path = self.root / "settings.yaml"
        raw = yaml.safe_load(path.read_text(encoding="utf-8")) if path.exists() else {}
        if raw is not None and not isinstance(raw, dict):
            raise ValueError("settings.yaml must contain a mapping.")
        return DEFAULT_SETTINGS | (raw or {})

    def save_settings(self, settings):
        cdp_url(settings["cdp_url"])
        if not 1 <= int(settings["timeout_seconds"]) <= 300:
            raise ValueError("Timeout must be 1–300 seconds.")
        atomic_text(self.root / "settings.yaml", yaml.safe_dump(settings, sort_keys=False))

    def save(self, workflow):
        workflow = validate_workflow(workflow)
        body = json.dumps(workflow, ensure_ascii=False, indent=2)
        atomic_text(self.workflows / (workflow["id"] + ".json"), body)
        with self.db:
            self.db.execute("INSERT OR REPLACE INTO workflows VALUES(?,?,?,?)", (workflow["id"], workflow["name"], body, now()))
        return workflow

    def list_workflows(self):
        return [json.loads(row[0]) for row in self.db.execute("SELECT body FROM workflows ORDER BY updated DESC, name")]

    def start(self, workflow, mode="running"):
        run_id = str(uuid4())
        with self.db:
            self.db.execute("INSERT INTO runs VALUES(?,?,?,?,?,?,?,?)", (run_id, workflow["id"], workflow["name"], now(), None, mode, "", ""))
        return run_id

    def event(self, run_id, event):
        with self.db:
            self.db.execute("INSERT INTO events(run_id,time,body) VALUES(?,?,?)", (run_id, now(), json.dumps(event, ensure_ascii=False)))

    def finish(self, run_id, status, artifact="", detail=""):
        with self.db:
            self.db.execute("UPDATE runs SET finished=?, status=?, artifact=?, detail=? WHERE id=?", (now(), status, artifact, detail, run_id))

    def runs(self):
        return [dict(r) for r in self.db.execute("SELECT * FROM runs ORDER BY started DESC, rowid DESC LIMIT 200")]

    def events(self, run_id):
        return [json.loads(r[0]) for r in self.db.execute("SELECT body FROM events WHERE run_id=? ORDER BY id", (run_id,))]
